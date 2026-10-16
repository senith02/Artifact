"""Build Word copies of the dissertation chapters and a figure gallery, for pasting into the thesis.

Not part of the research stack: needs `pip install pypandoc_binary python-docx pillow` in any Python
environment. Reads the chapter Markdown files and dissertation/figures/ (never modifies them) and
writes dissertation/word/:
  - chapter4_proposed_solution.docx, chapter5_evaluation.docx, chapter6_discussion.docx
    (chapter 6 is followed by the appendices from appendices.md, each on a new page)
  - figures_gallery.docx  (every figure with its caption, sized to fit an A4 text block)

Run from dissertation/:  python build_word.py
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
OUT = HERE / "word"
CHAPTERS = {
    "proposed solution.md": "chapter4_proposed_solution.docx",
    "chapter5_evaluation.md": "chapter5_evaluation.docx",
    "chapter6_discussion.md": "chapter6_discussion.docx",
}
# Extra Markdown files appended to a chapter's Word copy (the appendices follow Chapter 6; in the
# thesis they go after the References).
APPENDED = {"chapter6_discussion.md": ["appendices.md"]}
MAX_W_CM, MAX_H_CM = 16.0, 20.0  # A4 with 2.5 cm margins, leaving room for the caption

IMG = re.compile(r"!\[[^\]]*\]\((figures/[^)]+\.png)\)")
CAPTION = re.compile(r"^\*\*(Figure [A-Z\d]+\.\d+ — .+?)\*\*(.*)$")


def fit_cm(png: Path) -> tuple[float, float]:
    w, h = Image.open(png).size
    scale = min(MAX_W_CM / w, MAX_H_CM / h)
    return w * scale, h * scale


def chapter_docx(md_name: str, docx_name: str) -> None:
    sources = [md_name, *APPENDED.get(md_name, [])]
    text = "\n\n".join((HERE / name).read_text(encoding="utf-8") for name in sources)

    def sized(m: re.Match) -> str:
        wcm, _ = fit_cm(HERE / m.group(1))
        return f"![]({m.group(1)}){{width={wcm:.2f}cm}}"

    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8", dir=HERE) as tmp:
        tmp.write(IMG.sub(sized, text))
    try:
        pypandoc.convert_file(tmp.name, "docx", format="gfm+attributes",
                              outputfile=str(OUT / docx_name), extra_args=[f"--resource-path={HERE}"])
    finally:
        Path(tmp.name).unlink()
    if len(sources) > 1:
        # Each appendix starts on a new page.
        doc = Document(OUT / docx_name)
        for p in doc.paragraphs:
            if p.style.name == "Heading 1" and p.text.startswith("Appendi"):
                p.paragraph_format.page_break_before = True
        doc.save(OUT / docx_name)


def gallery() -> int:
    """Pair each '**Figure x.y — caption**' line with the next figure image in the chapter."""
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    doc.add_heading("Dissertation figures", level=1)
    doc.add_paragraph("Each figure below is the same 300-dpi PNG as in dissertation/figures/. "
                      "Copy a figure and its caption into the thesis, or use Insert → Pictures on the "
                      "PNG or SVG file directly.")
    n = 0
    for md_name in [name for md in CHAPTERS for name in (md, *APPENDED.get(md, []))]:
        lines = (HERE / md_name).read_text(encoding="utf-8").splitlines()
        pending = None
        for line in lines:
            c = CAPTION.match(line.strip())
            if c:
                pending = (c.group(1) + c.group(2).replace("*", "")).strip()
                continue
            m = IMG.search(line)
            if m and pending:
                png = HERE / m.group(1)
                wcm, hcm = fit_cm(png)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(png), width=Cm(wcm), height=Cm(hcm))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(pending)
                run.italic = True
                run.font.size = Pt(10)
                doc.add_page_break()
                n += 1
                pending = None
    doc.save(OUT / "figures_gallery.docx")
    return n


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for md_name, docx_name in CHAPTERS.items():
        chapter_docx(md_name, docx_name)
        print(f"wrote word/{docx_name}")
    print(f"wrote word/figures_gallery.docx ({gallery()} figures)")


if __name__ == "__main__":
    main()
